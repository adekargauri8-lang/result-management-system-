from openpyxl import Workbook, load_workbook
import os


FILE_NAME = "student_results.xlsx"


# -------------------------------
# Calculate Result
# -------------------------------
def calculate_result(marks):
    total = sum(marks)
    percentage = total / 5

    if percentage >= 75:
        grade = "A"
    elif percentage >= 60:
        grade = "B"
    elif percentage >= 50:
        grade = "C"
    elif percentage >= 40:
        grade = "D"
    else:
        grade = "F"

    if all(mark >= 35 for mark in marks):
        status = "PASS"
    else:
        status = "FAIL"
        grade = "F"

    return total, percentage, grade, status


# -------------------------------
# Create Excel File
# -------------------------------
def create_excel_file():
    if not os.path.exists(FILE_NAME):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Student Results"

        headers = [
            "Roll No", "Name", "Class",
            "Subject 1", "Subject 2", "Subject 3",
            "Subject 4", "Subject 5",
            "Total", "Percentage", "Grade", "Status"
        ]

        sheet.append(headers)
        workbook.save(FILE_NAME)


# -------------------------------
# Add Student
# -------------------------------
def add_student():
    create_excel_file()

    print("\n--------------------------------")
    print("       ADD STUDENT RESULT")
    print("--------------------------------")

    roll_no = input("Enter Roll No: ")
    name = input("Enter Student Name: ")
    student_class = input("Enter Course/Class: ")

    marks = []

    for i in range(1, 6):
        while True:
            try:
                mark = float(input(f"Enter Marks of Subject {i}: "))

                if 0 <= mark <= 100:
                    marks.append(mark)
                    break
                else:
                    print("Marks must be between 0 and 100.")

            except ValueError:
                print("Please enter a valid number.")

    total, percentage, grade, status = calculate_result(marks)

    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    sheet.append([
        roll_no,
        name,
        student_class,
        marks[0],
        marks[1],
        marks[2],
        marks[3],
        marks[4],
        total,
        percentage,
        grade,
        status
    ])

    workbook.save(FILE_NAME)

    print("\nStudent result added successfully!")
    print("--------------------------------")
    print(f"Total      : {total}")
    print(f"Percentage : {percentage:.2f}%")
    print(f"Grade      : {grade}")
    print(f"Status     : {status}")
    print("--------------------------------")


# -------------------------------
# Get Student Result
# -------------------------------
def get_result():
    create_excel_file()

    roll_no = input("\nEnter Roll No: ")

    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    found = False

    for row in sheet.iter_rows(min_row=2, values_only=True):

        if str(row[0]) == str(roll_no):

            print("\n--------------------------------")
            print("          STUDENT RESULT")
            print("--------------------------------")
            print(f"Roll No     : {row[0]}")
            print(f"Name        : {row[1]}")
            print(f"Class       : {row[2]}")
            print(f"Total       : {row[8]}")
            print(f"Percentage  : {row[9]:.2f}%")
            print(f"Grade       : {row[10]}")
            print(f"Status      : {row[11]}")
            print("--------------------------------")

            found = True
            break

    if not found:
        print("\nStudent with this Roll No. not found.")


# -------------------------------
# Show All Student Data
# -------------------------------
def show_all_data():
    create_excel_file()

    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    print("\n================ STUDENT DATA ================")

    print(
        f"{'Roll No':<10}"
        f"{'Name':<15}"
        f"{'Class':<10}"
        f"{'Total':<10}"
        f"{'Percentage':<12}"
        f"{'Grade':<8}"
        f"{'Status':<8}"
    )

    print("-" * 73)

    data_found = False

    for row in sheet.iter_rows(min_row=2, values_only=True):

        data_found = True

        print(
            f"{str(row[0]):<10}"
            f"{str(row[1]):<15}"
            f"{str(row[2]):<10}"
            f"{str(row[8]):<10}"
            f"{row[9]:<12.2f}"
            f"{str(row[10]):<8}"
            f"{str(row[11]):<8}"
        )

    if not data_found:
        print("No student records found.")

    print("=" * 73)


# -------------------------------
# Menu
# -------------------------------
def menu():

    create_excel_file()

    while True:

        print("\n========================================")
        print("      STUDENT RESULT MANAGEMENT")
        print("========================================")
        print("1. Add Student Result")
        print("2. Get Student Result")
        print("3. Show All Student Data")
        print("4. Exit")
        print("========================================")

        choice = input("Enter your choice: ")

        if choice == "1":
            add_student()

        elif choice == "2":
            get_result()

        elif choice == "3":
            show_all_data()

        elif choice == "4":
            print("\nThank you for using Student Result Management System.")
            break

        else:
            print("\nInvalid choice! Please enter 1 to 4.")


# -------------------------------
# Start Program
# -------------------------------
menu()